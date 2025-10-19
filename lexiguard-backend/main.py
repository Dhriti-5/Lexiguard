# main.py - FIXED VERSION WITH PROPER DLP INTEGRATION

import os
import json
import io
import logging
import time
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from dotenv import load_dotenv
import google.generativeai as genai
from google.cloud import dlp_v2
from google.cloud.dlp_v2 import types as dlp_types
import PyPDF2
from docx import Document

# --- 0. CONFIGURE LOGGING ---
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# --- 1. LOAD ENVIRONMENT VARIABLES ---
load_dotenv()

# --- 2. INITIALIZE FASTAPI APP ---
app = FastAPI(
    title="LexiGuard API",
    description="Analyzes legal documents (text, PDF, or DOCX) using Google's Gemini AI with PII Redaction.",
    version="1.4.0"
)

# --- 3. ENABLE CORS ---
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- 4. CONFIGURE GOOGLE GEMINI ---
API_KEY = os.getenv("GOOGLE_API_KEY")
if not API_KEY:
    logger.warning("GOOGLE_API_KEY not found in .env. Relying on ADC or Cloud Run credentials.")
else:
    genai.configure(api_key=API_KEY)

safety_settings = {
    "HARM_CATEGORY_HARASSMENT": "block_none",
    "HARM_CATEGORY_HATE_SPEECH": "block_none",
    "HARM_CATEGORY_SEXUALLY_EXPLICIT": "block_none",
    "HARM_CATEGORY_DANGEROUS_CONTENT": "block_none",
}

# Initialize Gemini model
MODEL_NAME = "models/gemini-2.5-flash"
model = None

try:
    model = genai.GenerativeModel(MODEL_NAME, safety_settings=safety_settings)
    logger.info(f"✅ Gemini model initialized successfully with {MODEL_NAME}")
except Exception as e:
    logger.error(f"❌ Failed to initialize Gemini model '{MODEL_NAME}': {e}")
    model = None

# --- 5. CONFIGURE GOOGLE CLOUD DLP ---
dlp_client = None
PROJECT_ID = os.getenv("GOOGLE_CLOUD_PROJECT")

def get_dlp_client():
    """Lazy initialization of DLP client"""
    global dlp_client
    if dlp_client is None:
        try:
            dlp_client = dlp_v2.DlpServiceClient()
            logger.info("✅ DLP client initialized successfully")
        except Exception as e:
            logger.warning(f"⚠️ DLP client initialization failed: {e}")
            logger.warning("PII redaction will be disabled")
    return dlp_client

if not PROJECT_ID:
    logger.warning("GOOGLE_CLOUD_PROJECT not set; DLP may fail without ADC context.")

# Define info types to redact
INFO_TYPES_TO_REDACT = [
    dlp_types.InfoType(name="PERSON_NAME"),
    dlp_types.InfoType(name="EMAIL_ADDRESS"),
    dlp_types.InfoType(name="PHONE_NUMBER"),
    dlp_types.InfoType(name="STREET_ADDRESS"),
    dlp_types.InfoType(name="CREDIT_CARD_NUMBER"),
    dlp_types.InfoType(name="DATE_OF_BIRTH"),
    dlp_types.InfoType(name="US_SOCIAL_SECURITY_NUMBER"),
]

info_type_transformations = dlp_types.InfoTypeTransformations(
    transformations=[
        {
            "info_types": INFO_TYPES_TO_REDACT,
            "primitive_transformation": dlp_types.PrimitiveTransformation(
                replace_with_info_type_config=dlp_types.ReplaceWithInfoTypeConfig()
            ),
        }
    ]
)

DEIDENTIFY_CONFIG = dlp_types.DeidentifyConfig(
    info_type_transformations=info_type_transformations
)

# --- 6. DATA MODELS ---
class DocumentRequest(BaseModel):
    text: str

class ChatRequest(BaseModel):
    message: str
    document_text: str

class NegotiationRequest(BaseModel):
    clause: str

class DocumentEmailRequest(BaseModel):
    document_summary: str
    risk_summary: str

class ExtendedAnalysisRequest(BaseModel):
    text: str

# --- 7. PROMPTS ---
SUMMARY_PROMPT = """
You are LexiGuard, an expert AI assistant that explains complex legal documents in simple terms.
Analyze the following contract. Provide a concise, bullet-point summary covering:
1. The primary purpose of the agreement.
2. The key responsibilities of each party.
3. The duration and key financial terms (like rent, salary, etc.).
Use clear, simple language suitable for a non-lawyer.

Note: Personal information has been replaced with placeholders like [PERSON_NAME], [EMAIL_ADDRESS], etc.
"""

RISK_ANALYSIS_PROMPT = """
You are a meticulous risk analysis AI. Scan the provided legal document.
Your task is to identify and extract any clauses that are potentially unfavorable, non-standard, or represent a significant risk.
Focus specifically on clauses related to: Indemnity, Limitation of Liability, Automatic Renewal, Termination Penalties, and Non-Compete agreements.
For each identified clause, you MUST provide the original text, a simple one-sentence explanation of the risk, and a severity level of either 'High' or 'Medium'.
You MUST respond ONLY with a valid JSON object. The structure of the JSON object must be:
{"risks": [{"clause_text": "...", "risk_explanation": "...", "severity": "..."}]}
If you find no risks, you MUST return: {"risks": []}
Do not add any text or formatting before or after the JSON object.

Note: Personal information has been replaced with placeholders like [PERSON_NAME], [EMAIL_ADDRESS], etc.
"""

NEGOTIATION_PROMPT = """
You are LexiGuard, an AI that helps users politely negotiate risky contract clauses.
Draft a professional and polite email body requesting to amend or clarify the following clause.

Clause (PII has been redacted with placeholders like [PERSON_NAME], [EMAIL_ADDRESS]):
{clause}

Your task is to draft ONLY the email body. The email should:
1. Start professionally (e.g., "Dear [Recipient Name],").
2. Clearly reference the clause in question.
3. Explain the user's concern politely.
4. Suggest a discussion or a more balanced alternative.
5. Use standard placeholders like [Your Name], [Your Company], [Recipient Name], [Date] where needed.
6. End with a collaborative closing.

Generate ONLY the email body text.
"""

DOCUMENT_EMAIL_PROMPT = """
You are LexiGuard, an AI assistant that helps users communicate about legal document reviews.

Generate a professional email to send to a legal advisor, counterparty, or stakeholder regarding a legal document review.

Document Summary (PII redacted):
{document_summary}

Identified Risks (PII redacted):
{risk_summary}

The email should:
1. Be professional, clear, and concise
2. Summarize the key findings from the document analysis
3. Highlight the most critical risks identified
4. Request clarification, revision, or discussion on the concerning clauses
5. Maintain a collaborative and constructive tone
6. Use placeholders like [Your Name], [Company Name], [Recipient Name] where appropriate

Generate ONLY the email body. Do not include subject line.
Start directly with a professional greeting.
"""

FAIRNESS_PROMPT = """
You are LexiGuard, a fairness evaluator.
Compare the following risky clause with a standard, balanced contract clause.
Return a JSON object strictly in this format:
{
  "standard_clause": "...",
  "risky_clause": "...",
  "fairness_score": 0-100,
  "explanation": "..."
}

Risky Clause (PII may be redacted):
{clause}
"""

DETAILED_CLAUSE_ANALYSIS_PROMPT = """
You are a legal expert analyzing contracts and agreements for LexiGuard. 
Analyze the following document and identify ALL risky or concerning clauses with deep explanations.

For EACH risky clause, provide:
1. **clause**: The exact clause text or relevant excerpt (keep it under 200 characters if too long)
2. **risk_level**: "High", "Medium", or "Low"
3. **impact**: Brief description of potential harm to the user (1 sentence)
4. **recommendation**: Specific actionable advice for negotiation (1-2 sentences)
5. **explanation**: Detailed plain-language explanation of why this clause is risky and what could go wrong (2-3 sentences)

Focus on these risk categories:
- Termination rights (sudden eviction, firing without cause)
- Financial liability (unlimited damages, penalties)
- Automatic renewal traps
- Non-compete restrictions
- Indemnification clauses
- Limitation of liability
- Unilateral changes by one party
- Waiver of legal rights

You MUST respond ONLY with a valid JSON array. Use this exact format:
[
    {
        "clause": "Original clause text here",
        "risk_level": "High",
        "impact": "Brief impact description",
        "recommendation": "What the user should do or negotiate",
        "explanation": "Detailed plain-language explanation of why this is risky"
    }
]

If no risks are found, return an empty array: []
CRITICAL: Respond ONLY with valid JSON, no additional text, no markdown.

Note: Personal information has been replaced with placeholders like [PERSON_NAME], [EMAIL_ADDRESS], etc.
"""

CHAT_PROMPT = """
You are LexiGuard, a helpful AI assistant specializing in legal document analysis.

The document you're discussing has had Personal Identifiable Information (PII) replaced with placeholders:
- [PERSON_NAME] - replaced names
- [EMAIL_ADDRESS] - replaced emails
- [PHONE_NUMBER] - replaced phone numbers
- [STREET_ADDRESS] - replaced addresses
- [DATE_OF_BIRTH] - replaced dates of birth
- etc.

When answering questions:
1. Focus on the legal terms, obligations, and structure of the document
2. If asked about specific names/addresses that are redacted, acknowledge that "personal information has been redacted for privacy"
3. Explain what the clause or section DOES, regardless of who it applies to
4. Be clear and concise
5. Use non-legal language that anyone can understand

Document Context (with PII redacted):
{document_text}

User Question:
{message}

Provide a helpful answer based on the document content.
"""

# --- 8. HELPER FUNCTIONS ---

def redact_text_with_dlp(text: str):
    """Redact PII from text using Google Cloud DLP"""
    if not text or not PROJECT_ID:
        logger.warning("DLP: Project ID not configured or text is empty. Skipping redaction.")
        return text, False

    client = get_dlp_client()
    if not client:
        logger.warning("DLP client not available. Skipping redaction.")
        return text, False

    parent_path = f"projects/{PROJECT_ID}/locations/global"
    item = {"value": text}

    try:
        response = client.deidentify_content(
            request={
                "parent": parent_path,
                "deidentify_config": DEIDENTIFY_CONFIG,
                "inspect_config": {"info_types": INFO_TYPES_TO_REDACT},
                "item": item,
            }
        )
        redacted = response.item.value
        changed = redacted != text
        if changed:
            logger.info("✅ DLP Redaction complete - PII found and redacted")
        else:
            logger.info("ℹ️ DLP Redaction complete - No PII found")
        return redacted, changed
    except Exception as e:
        logger.error(f"❌ DLP failed: {e}")
        return text, False

def extract_text_from_pdf(file_stream):
    """Extract text from PDF file"""
    try:
        pdf = PyPDF2.PdfReader(file_stream)
        return "".join([p.extract_text() or "" for p in pdf.pages])
    except Exception as e:
        logger.error(f"PDF extraction failed: {e}")
        raise HTTPException(status_code=500, detail="Error extracting text from PDF.")

def extract_text_from_docx(file_stream):
    """Extract text from DOCX file"""
    try:
        doc = Document(file_stream)
        return "\n".join([p.text for p in doc.paragraphs])
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        raise HTTPException(status_code=500, detail="Error extracting text from DOCX.")

def extract_text_from_txt(file_stream):
    """Extract text from TXT file"""
    try:
        content = file_stream.read()
        if isinstance(content, bytes):
            text = content.decode('utf-8')
        else:
            text = content
        return text
    except UnicodeDecodeError:
        file_stream.seek(0)
        content = file_stream.read()
        return content.decode('latin-1')

# --- 9. CORE ANALYSIS LOGIC ---

def analyze_text_internal(text: str):
    """Analyze text with PII redaction"""
    if model is None:
        raise HTTPException(status_code=503, detail="AI model not initialized. Check backend logs.")
    
    try:
        # Redact PII first
        redacted_text, pii_found = redact_text_with_dlp(text)
        
        # Generate summary
        summary_prompt = f"{SUMMARY_PROMPT}\n\nDocument:\n{redacted_text}"
        summary_response = model.generate_content(summary_prompt)
        summary = summary_response.text.strip()

        # Analyze risks
        risk_prompt = f"{RISK_ANALYSIS_PROMPT}\n\nDocument:\n{redacted_text}"
        risk_response = model.generate_content(risk_prompt)
        
        try:
            risks_text = risk_response.text.strip().replace("```json", "").replace("```", "").strip()
            risks_data = json.loads(risks_text)
        except Exception as e:
            logger.error(f"Risk JSON parse error: {e}")
            risks_data = {"risks": []}

        return {
            "summary": summary,
            "risks": risks_data.get("risks", []),
            "pii_redacted": pii_found,
            "redacted_text": redacted_text
        }
    except Exception as e:
        logger.error(f"Error in analyze_text_internal: {e}")
        raise HTTPException(status_code=500, detail=f"Analysis failed: {str(e)}")

def analyze_clauses_detailed_internal(text: str):
    """Detailed clause analysis with PII redaction"""
    if model is None:
        raise HTTPException(status_code=503, detail="AI model not initialized. Check backend logs.")
    
    try:
        # Redact PII first
        redacted_text, pii_found = redact_text_with_dlp(text)
        
        prompt = f"{DETAILED_CLAUSE_ANALYSIS_PROMPT}\n\nDocument:\n{redacted_text}"
        response = model.generate_content(prompt)
        
        try:
            risks_text = response.text.strip().replace("```json", "").replace("```", "").strip()
            clauses = json.loads(risks_text)
        except Exception as e:
            logger.error(f"Clause JSON parse error: {e}")
            clauses = []
        
        return {
            "clauses": clauses if isinstance(clauses, list) else [],
            "pii_redacted": pii_found,
            "redacted_text": redacted_text
        }
    except Exception as e:
        logger.error(f"Error in detailed clause analysis: {e}")
        return {"clauses": [], "pii_redacted": False, "redacted_text": text}

# --- 10. ROUTES ---

@app.get("/")
async def root():
    return {
        "message": "LexiGuard API is running with PII Protection 🔒",
        "version": "1.4.0",
        "features": ["PII Redaction", "AI Analysis", "Risk Detection"],
        "endpoints": [
            "/analyze",
            "/analyze-file",
            "/analyze-clauses",
            "/draft-negotiation",
            "/draft-document-email",
            "/analyze-extended",
            "/chat"
        ],
        "supported_formats": ["PDF", "DOCX", "TXT", "Plain Text"]
    }

@app.post("/analyze")
async def analyze_document(request: DocumentRequest):
    """Analyze text document with PII redaction"""
    result = analyze_text_internal(request.text)
    
    response = {
        "file_type": "Text",
        "summary": result.get("summary", ""),
        "risks": result.get("risks", []),
        "suggestions": [],  # Generate from risks if needed
        "pii_redacted": result.get("pii_redacted", False)
    }
    
    if result.get("pii_redacted"):
        response["privacy_notice"] = "✓ Personal information has been redacted for your privacy"
    
    return response

@app.post("/analyze-file")
async def analyze_file(file: UploadFile = File(None), text: str = Form(None)):
    """Standard analysis with file upload support"""
    document_text = ""
    file_type = "Text"
    filename_display = "Direct Text Input"

    if file:
        filename_display = file.filename
        filename_lower = file.filename.lower()
        
        if filename_lower.endswith(".pdf"):
            document_text = extract_text_from_pdf(file.file)
            file_type = "PDF"
        elif filename_lower.endswith(".docx"):
            document_text = extract_text_from_docx(file.file)
            file_type = "DOCX"
        elif filename_lower.endswith(".txt"):
            document_text = extract_text_from_txt(file.file)
            file_type = "TXT"
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type. Only PDF, DOCX, and TXT allowed.")
    elif text:
        document_text = text
    else:
        raise HTTPException(status_code=400, detail="No file or text provided")

    if not document_text.strip():
        raise HTTPException(status_code=400, detail="No text could be extracted from the document")

    result = analyze_text_internal(document_text)

    response = {
        "filename": filename_display,
        "file_type": file_type,
        "summary": result.get("summary", ""),
        "risks": result.get("risks", []),
        "suggestions": [],
        "pii_redacted": result.get("pii_redacted", False)
    }

    if result.get("pii_redacted"):
        response["privacy_notice"] = "✓ Personal information has been redacted for your privacy"

    return response

@app.post("/analyze-clauses")
async def analyze_clauses(file: UploadFile = File(None), text: str = Form(None)):
    """Detailed clause analysis with file upload support"""
    document_text = ""
    file_type = "Text"
    filename_display = "Direct Text Input"

    if file:
        filename_display = file.filename
        filename_lower = file.filename.lower()
        
        if filename_lower.endswith(".pdf"):
            document_text = extract_text_from_pdf(file.file)
            file_type = "PDF"
        elif filename_lower.endswith(".docx"):
            document_text = extract_text_from_docx(file.file)
            file_type = "DOCX"
        elif filename_lower.endswith(".txt"):
            document_text = extract_text_from_txt(file.file)
            file_type = "TXT"
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type. Only PDF, DOCX, and TXT allowed.")
    elif text:
        document_text = text
    else:
        raise HTTPException(status_code=400, detail="No file or text provided")

    if not document_text.strip():
        raise HTTPException(status_code=400, detail="No text could be extracted from the document")

    result = analyze_clauses_detailed_internal(document_text)

    response = {
        "filename": filename_display,
        "file_type": file_type,
        "total_risky_clauses": len(result.get("clauses", [])),
        "clauses": result.get("clauses", []),
        "document_preview": result.get("redacted_text", document_text)[:300],
        "pii_redacted": result.get("pii_redacted", False)
    }

    if result.get("pii_redacted"):
        response["privacy_notice"] = "✓ Personal information has been redacted for your privacy"

    return response

@app.post("/draft-negotiation")
async def draft_negotiation(request: NegotiationRequest):
    """Generate negotiation email for a risky clause"""
    if model is None:
        raise HTTPException(status_code=503, detail="AI model not initialized")
    
    # Redact PII from the clause first
    redacted_clause, _ = redact_text_with_dlp(request.clause)
    
    prompt = NEGOTIATION_PROMPT.format(clause=redacted_clause)
    
    try:
        response = model.generate_content(prompt)
        email_text = response.text.strip() or "Could not generate email."
    except Exception as e:
        logger.error(f"Negotiation email generation error: {e}")
        email_text = f"Error: {str(e)}"
    
    return {"negotiation_email": email_text}

@app.post("/draft-document-email")
async def draft_document_email(request: DocumentEmailRequest):
    """Generate comprehensive document review email"""
    if model is None:
        raise HTTPException(status_code=503, detail="AI model not initialized")
    
    # Redact PII from summary and risks
    redacted_summary, _ = redact_text_with_dlp(request.document_summary)
    redacted_risks, _ = redact_text_with_dlp(request.risk_summary)
    
    prompt = DOCUMENT_EMAIL_PROMPT.format(
        document_summary=redacted_summary[:2000],
        risk_summary=redacted_risks[:2000]
    )
    
    try:
        response = model.generate_content(prompt)
        email_text = response.text.strip() or "Could not generate email."
    except Exception as e:
        logger.error(f"Document email generation error: {e}")
        email_text = f"Error: {str(e)}"
    
    return {"document_email": email_text}

@app.post("/analyze-extended")
async def analyze_extended(request: ExtendedAnalysisRequest):
    """Performs extended analysis with clause comparison and fairness scoring"""
    if model is None:
        raise HTTPException(status_code=503, detail="AI model not initialized")
    
    base_result = analyze_text_internal(request.text)

    fairness_results = []
    for risk in base_result.get("risks", []):
        clause_text = risk.get("clause_text", "")
        if not clause_text.strip():
            continue
        
        # Redact PII from clause
        redacted_clause, _ = redact_text_with_dlp(clause_text)
        
        prompt = FAIRNESS_PROMPT.format(clause=redacted_clause)
        try:
            resp = model.generate_content(prompt)
            fairness_json = resp.text.strip().replace("```json", "").replace("```", "").strip()
            parsed = json.loads(fairness_json)
        except Exception as e:
            logger.error(f"Fairness analysis error: {e}")
            parsed = {
                "standard_clause": "",
                "risky_clause": redacted_clause,
                "fairness_score": 50,
                "explanation": f"Error: {str(e)}"
            }
        fairness_results.append(parsed)
        time.sleep(2)  # Rate limiting

    return {
        "summary": base_result.get("summary", ""),
        "risks": base_result.get("risks", []),
        "fairness_analysis": fairness_results,
        "pii_redacted": base_result.get("pii_redacted", False)
    }

@app.post("/chat")
async def chat_with_document(request: ChatRequest):
    """Chat with document - handles redacted PII properly"""
    if model is None:
        raise HTTPException(status_code=503, detail="AI model not initialized")
    
    if not request.message.strip() or not request.document_text.strip():
        return {"reply": "Please provide both a message and document text."}

    # Redact PII from the document text that's passed to chat
    redacted_document, pii_found = redact_text_with_dlp(request.document_text)
    
    # Also redact the user's question in case they mention PII
    redacted_message, _ = redact_text_with_dlp(request.message)

    prompt = CHAT_PROMPT.format(
        document_text=redacted_document,
        message=redacted_message
    )
    
    try:
        response = model.generate_content(prompt)
        answer = response.text.strip() or "No answer could be generated."
    except Exception as e:
        logger.error(f"Chat error: {e}")
        answer = f"Error: {str(e)}"
    
    return {
        "reply": answer,
        "pii_redacted": pii_found
    }

# --- RUN SERVER ---
if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run("main:app", host="0.0.0.0", port=port, reload=True)